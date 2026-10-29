"""Tests de plugin_tools: el asistente del chat usando Scraping, SEO y Office.

Los plugins premium no viajan con el navegador, así que aquí se simulan sus
paneles e integración con la misma interfaz que usa main_window.
"""
import os
import sys

import pytest
from PySide6.QtWidgets import QApplication

import plugin_tools as PT

VIEWER_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                          "plugins", "office_viewer")
if os.path.isdir(VIEWER_DIR) and VIEWER_DIR not in sys.path:
    sys.path.insert(0, VIEWER_DIR)


@pytest.fixture(scope="session")
def qapp():
    return QApplication.instance() or QApplication([])


# ── Funciones puras ─────────────────────────────────────────────────────────

def test_cell_value_numbers_and_text():
    assert PT.cell_value("12") == 12
    assert PT.cell_value("12.5") == 12.5
    assert PT.cell_value("1.234,56") == 1234.56
    assert PT.cell_value("1.234") == 1234
    assert PT.cell_value("-3,5") == -3.5
    assert PT.cell_value("12,50 €") == "12,50 €"
    assert PT.cell_value("A1") == "A1"
    assert PT.cell_value("") == ""


def test_parse_markdown_table():
    title, rows = PT.parse_markdown_table(
        "Comparativa\n| Modelo | Precio |\n|---|:---:|\n| A | 999 |\n| B |\n")
    assert title == "Comparativa"
    assert rows == [["Modelo", "Precio"], ["A", "999"], ["B", ""]]


def test_parse_markdown_table_without_table():
    assert PT.parse_markdown_table("solo texto") == ("solo texto", [])


def test_effective_pattern_selector():
    generic = {"selectors": {"recommended": "div", "common_ancestor": "section.grid"}}
    assert PT.effective_pattern_selector(generic) == "section.grid > div"
    assert PT.effective_pattern_selector({"selectors": {"recommended": "div.card"}}) == "div.card"
    assert PT.effective_pattern_selector({"selectors": {"index_selectors": ["ul > li:nth-child(1)"]}}) \
        == "ul > li:nth-child(1)"
    assert PT.effective_pattern_selector({}) == ""


# ── Simulación del navegador y los plugins ──────────────────────────────────

class FakeUrl:
    def __init__(self, url):
        self._url = url

    def toString(self):
        return self._url

    def scheme(self):
        return self._url.split(":", 1)[0]

    def isEmpty(self):
        return not self._url


class FakePage:
    def __init__(self, html):
        self.html = html

    def toHtml(self, cb):
        cb(self.html)


class FakeBrowser:
    def __init__(self, url, html=""):
        self._url = FakeUrl(url)
        self._page = FakePage(html)

    def url(self):
        return self._url

    def page(self):
        return self._page


class FakeTabs:
    def __init__(self, browser):
        self.browser = browser

    def currentWidget(self):
        return self.browser


class FakeTabManager:
    def __init__(self, browser):
        self.tabs = FakeTabs(browser)


class FakePluginManager:
    def __init__(self, ids):
        self.plugins = {i: object() for i in ids}


class FakeIntegration:
    EXPORT_COLUMNS = ["n", "titulo", "precio", "texto", "enlace"]

    def __init__(self):
        self.current_html = ""
        self.current_url = ""
        self.selected = []
        self.patterns = [
            {"type": "cards", "elements_count": 3, "similarity_score": 0.9,
             "selectors": {"recommended": "div.product"},
             "sample_items": [{"text": "Producto 1 9,99 €"}]},
            {"type": "lists", "items_count": 5, "selectors": {"recommended": ""}},
        ]

    def update_content(self, html, url=""):
        self.current_html, self.current_url = html, url
        return True

    def detect_dom_patterns(self):
        return {"success": True, "patterns": self.patterns}

    def add_element_by_selector(self, selector):
        if selector != "div.product":
            return {"error": f"No elements found with selector: {selector}"}
        for i in range(1, 4):
            self.selected.append({"product_title": f"Producto {i}", "product_price": f"{i},99 €"})
        return {"success": True, "elements_added": 3, "selector": selector}

    def get_selected_elements(self):
        return self.selected

    def clear_selected_elements(self):
        self.selected.clear()

    def _build_export_rows(self):
        return [{"n": i, "titulo": e["product_title"], "precio": e["product_price"]}
                for i, e in enumerate(self.selected, 1)]


class FakeScrapingPanel:
    def __init__(self):
        self.updated = 0
        self.exported = None

        class Combo:
            text = ""

            def setCurrentText(self, t):
                Combo.text = t
        self.export_format = Combo()

    def update_selected_list(self):
        self.updated += 1

    def display_patterns(self, patterns):
        self.shown_patterns = patterns

    def run_export(self):
        self.exported = self.export_format.text


class FakeCoordinator:
    def analyze(self, html, url):
        return {"meta": {"score": 60, "severity": "warning",
                         "issues": [{"type": "warning", "message": "Falta meta description"}]},
                "overall": {"score": 60, "total_analyzers": 1}}

    def get_summary(self):
        return {"critical_issues": 0, "warnings": 1, "info": 0}

    def get_analyzer_info(self, name):
        return {"name": "Meta Tags"}

    def get_recommendations(self):
        return ["Añade una meta description"]


class FakeSeoPanel:
    def __init__(self):
        self.coordinator = FakeCoordinator()
        self.current_results = {}
        self.current_keywords = []
        self.current_url = ""
        self.current_html = ""
        self.exported = None

        class Input:
            def setText(self, t):
                pass
        self.url_input = Input()

    def on_analysis_finished(self, results):
        self.current_results = results
        self.current_keywords = [{"keyword": "gatos", "type": "palabra", "count": 7,
                                  "density": 2.1, "in_key_tags": True}]

    def export_results(self, fmt):
        self.exported = fmt


class FakeMain:
    def __init__(self, plugins, url="https://tienda.example/", html="<html></html>"):
        self.plugin_manager = FakePluginManager(plugins)
        self.tab_manager = FakeTabManager(FakeBrowser(url, html))
        self.scraping_integration = FakeIntegration()
        self.scraping_panel = FakeScrapingPanel()
        self.seo_panel = FakeSeoPanel()


def make_tools(plugins, **kw):
    main = FakeMain(plugins, **kw)
    tools = PT.PluginTools(lambda: main)
    tools._show_panel = lambda *a: None
    return tools, main


def run(tools, tag, arg, qapp=None):
    out = []
    tools.run(tag, arg, lambda ok, msg, fb: out.append((ok, msg, fb)))
    if qapp is not None:
        for _ in range(200):
            if out:
                break
            qapp.processEvents()
            for t in list(tools._threads):
                t.wait(50)
    assert out, "la herramienta no terminó"
    return out[0]


# ── Disponibilidad y licencia ───────────────────────────────────────────────

def test_only_loaded_plugins_are_offered():
    tools, _ = make_tools(["seo_analyzer"])
    assert tools.available() == ["seo"]
    prompt = tools.prompt_section()
    assert "<seo>analyze</seo>" in prompt
    assert "<scrape>" not in prompt and "<office>" not in prompt


def test_no_plugins_no_prompt():
    tools, _ = make_tools([])
    assert tools.prompt_section() == ""


def test_unlicensed_plugin_is_refused():
    tools, _ = make_tools(["seo_analyzer"])
    ok, msg, feedback = run(tools, "scrape", "patterns")
    assert not ok and "licencia" in msg and "Plugin Store" in feedback


def test_unknown_or_internal_command_is_refused():
    tools, _ = make_tools(["seo_analyzer", "scraping"])
    for tag, arg in (("seo", "panel"), ("seo", "summary_text"), ("scrape", "load"), ("seo", "borra")):
        ok, msg, feedback = run(tools, tag, arg)
        assert not ok and "no existe" in feedback


# ── Scraping ────────────────────────────────────────────────────────────────

def test_scrape_patterns_then_pattern_selection():
    tools, main = make_tools(["scraping"])
    ok, msg, feedback = run(tools, "scrape", "patterns")
    assert ok and "1 patrones" in msg
    assert "div.product" in feedback and "Producto 1" in feedback
    assert main.scraping_integration.current_url == "https://tienda.example/"
    ok, msg, feedback = run(tools, "scrape", "pattern 1")
    assert ok and "3 elementos" in msg
    assert "Producto 3" in feedback and "3,99 €" in feedback
    assert main.scraping_panel.updated == 1


def test_scrape_pattern_out_of_range():
    tools, _ = make_tools(["scraping"])
    ok, _, feedback = run(tools, "scrape", "pattern 4")
    assert not ok and "patterns" in feedback


def test_scrape_select_bad_selector():
    tools, _ = make_tools(["scraping"])
    ok, _, feedback = run(tools, "scrape", "select .nada")
    assert not ok and "No elements found" in feedback


def test_scrape_needs_a_web_page():
    tools, _ = make_tools(["scraping"], url="katphi-app://office/viewer.html")
    ok, msg, _ = run(tools, "scrape", "patterns")
    assert not ok and "página web" in msg


def test_scrape_export_uses_panel_dialog():
    tools, main = make_tools(["scraping"])
    ok, _, _ = run(tools, "scrape", "export excel")
    assert not ok   # sin selección
    run(tools, "scrape", "select div.product")
    ok, msg, feedback = run(tools, "scrape", "export excel")
    assert ok and feedback == "" and main.scraping_panel.exported == "Excel"


def test_scrape_clear():
    tools, main = make_tools(["scraping"])
    run(tools, "scrape", "select div.product")
    ok, _, _ = run(tools, "scrape", "clear")
    assert ok and main.scraping_integration.selected == []


# ── SEO ─────────────────────────────────────────────────────────────────────

def test_seo_analyze(qapp):
    tools, main = make_tools(["seo_analyzer"])
    ok, msg, feedback = run(tools, "seo", "analyze", qapp)
    assert ok and "60/100" in msg
    assert "Falta meta description" in feedback and "gatos" in feedback
    assert main.seo_panel.current_url == "https://tienda.example/"


def test_seo_keywords_reuses_analysis(qapp):
    tools, main = make_tools(["seo_analyzer"])
    run(tools, "seo", "analyze", qapp)
    main.seo_panel.coordinator = None   # si volviera a analizar, fallaría
    ok, _, feedback = run(tools, "seo", "keywords")
    assert ok and "gatos — 7" in feedback


def test_seo_export_requires_analysis(qapp):
    tools, main = make_tools(["seo_analyzer"])
    ok, _, _ = run(tools, "seo", "export pdf")
    assert not ok
    run(tools, "seo", "analyze", qapp)
    ok, _, feedback = run(tools, "seo", "export pdf")
    assert ok and feedback == "" and main.seo_panel.exported == "pdf"


# ── Office ──────────────────────────────────────────────────────────────────

@pytest.fixture
def office_tools(tmp_path, monkeypatch):
    def factory(plugins):
        tools, main = make_tools(plugins)
        monkeypatch.setattr(PT.PluginTools, "_output_dir", staticmethod(lambda: str(tmp_path)))
        opened = []
        tools._open_in_office = lambda path: opened.append(path) or True
        return tools, main, opened
    return factory


def test_office_sheet_from_markdown(office_tools):
    import openpyxl
    tools, _, opened = office_tools(["office_suite"])
    ok, msg, _ = run(tools, "office", "sheet Comparativa portátiles\n| Modelo | Precio |\n|---|---|\n"
                                      "| A | 999,5 |\n| B | 1.200 |")
    assert ok and "2 filas" in msg
    path = opened[0]
    assert os.path.basename(path).startswith("katphi_comparativa_port")
    ws = openpyxl.load_workbook(path).active
    assert [c.value for c in ws[1]] == ["Modelo", "Precio"]
    assert ws["B2"].value == 999.5 and ws["B3"].value == 1200
    assert ws["A1"].font.bold and ws.freeze_panes == "A2"


def test_office_sheet_never_overwrites(office_tools):
    tools, _, opened = office_tools(["office_suite"])
    table = "sheet X\n| a |\n|---|\n| 1 |"
    run(tools, "office", table)
    run(tools, "office", table)
    assert len(set(opened)) == 2


def test_office_sheet_from_scraping(office_tools):
    import openpyxl
    tools, _, opened = office_tools(["office_suite", "scraping"])
    ok, _, _ = run(tools, "office", "sheet scraping")
    assert not ok   # sin selección
    run(tools, "scrape", "select div.product")
    ok, _, _ = run(tools, "office", "sheet scraping")
    ws = openpyxl.load_workbook(opened[0]).active
    assert ok and ws["B4"].value == "Producto 3"


def test_office_sheet_scraping_needs_scraping_plugin(office_tools):
    tools, _, _ = office_tools(["office_suite"])
    ok, msg, _ = run(tools, "office", "sheet scraping")
    assert not ok and "Scraping" in msg


def test_office_sheet_from_seo(office_tools, qapp):
    import openpyxl
    tools, _, opened = office_tools(["office_suite", "seo_analyzer"])
    run(tools, "seo", "analyze", qapp)
    ok, _, _ = run(tools, "office", "sheet seo")
    wb = openpyxl.load_workbook(opened[0])
    assert ok and wb.sheetnames == ["Resumen", "Problemas", "Palabras clave"]
    assert wb["Problemas"]["C2"].value == "Falta meta description"


def test_office_sheet_without_table(office_tools):
    tools, _, opened = office_tools(["office_suite"])
    ok, _, feedback = run(tools, "office", "sheet Algo")
    assert not ok and "markdown" in feedback and not opened


def test_office_read_requires_office_tab():
    tools, _ = make_tools(["office_suite"])
    pytest.importorskip("katphi_office")
    ok, msg, _ = run(tools, "office", "read")
    assert not ok and "no es un documento" in msg


def test_document_text_sheets():
    import json
    from types import SimpleNamespace
    model = {"sheets": [{"name": "Datos", "rows": [[{"v": "a"}, None, {"v": "b"}], [{"v": "1"}]]},
                        {"name": "Oculta", "hidden": True, "rows": [[{"v": "secreto"}]]}]}
    s = SimpleNamespace(family="sheets", model_json=json.dumps(model).encode())
    text = PT.PluginTools._document_text(s)
    assert "Hoja «Datos»" in text and "a\t\tb" in text and "secreto" not in text


# ── Informe SEO (plugin 1.1.0: panel con save_report) ──────────────────────

class FakeSeoPanelReports(FakeSeoPanel):
    def __init__(self):
        super().__init__()
        self.saved = []

    def save_report(self, fmt, path):
        with open(path, "w") as f:
            f.write(fmt)
        self.saved.append((fmt, path))
        return path


@pytest.fixture
def report_tools(tmp_path, monkeypatch):
    tools, main = make_tools(["seo_analyzer"])
    main.seo_panel = FakeSeoPanelReports()
    monkeypatch.setattr(PT.PluginTools, "_output_dir", staticmethod(lambda: str(tmp_path)))
    opened = []
    tools._open_file = lambda path, fmt: opened.append((fmt, path)) or True
    return tools, main, opened


def test_seo_report_analyzes_first_when_needed(report_tools, qapp):
    tools, main, opened = report_tools
    ok, msg, feedback = run(tools, "seo", "report pdf", qapp)
    assert ok and feedback == "" and "abierto" in msg
    assert main.seo_panel.current_results     # se analizó antes del informe
    fmt, path = main.seo_panel.saved[0]
    assert fmt == "pdf" and os.path.basename(path).startswith("informe_seo_tienda.example_")
    assert opened == [("pdf", path)]


def test_seo_report_reuses_analysis_and_aliases(report_tools, qapp):
    tools, main, opened = report_tools
    run(tools, "seo", "analyze", qapp)
    main.seo_panel.coordinator = None          # si volviera a analizar, fallaría
    for alias, fmt in (("word", "docx"), ("Excel", "xlsx"), ("markdown", "md"), ("html", "html")):
        ok, _, _ = run(tools, "seo", f"report {alias}")
        assert ok and main.seo_panel.saved[-1][0] == fmt
    paths = [p for _, p in main.seo_panel.saved]
    assert len(set(paths)) == len(paths)       # nunca sobrescribe


def test_seo_export_is_alias_of_report(report_tools, qapp):
    tools, main, _ = report_tools
    ok, _, _ = run(tools, "seo", "export pdf", qapp)
    assert ok and main.seo_panel.saved and main.seo_panel.exported is None   # sin diálogo


def test_seo_report_invalid_format(report_tools):
    tools, _, _ = report_tools
    ok, _, feedback = run(tools, "seo", "report pptx")
    assert not ok and "docx" in feedback


def test_seo_report_error_is_reported(report_tools, qapp):
    tools, main, _ = report_tools

    def boom(fmt, path):
        raise PermissionError("El formato DOCX no está disponible en tu plan")
    main.seo_panel.save_report = boom
    ok, msg, feedback = run(tools, "seo", "report docx", qapp)
    assert not ok and "plan" in msg and "plan" in feedback


def test_seo_analysis_feedback_mentions_report(qapp):
    tools, _ = make_tools(["seo_analyzer"])
    _, _, feedback = run(tools, "seo", "analyze", qapp)
    assert "<seo>report FORMATO</seo>" in feedback


def test_prompt_mentions_report_formats():
    tools, _ = make_tools(["seo_analyzer"])
    assert "<seo>report pdf</seo>" in tools.prompt_section()
