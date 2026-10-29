"""Documentos de prueba para Katphi Office, generados al vuelo."""
import io
import os
import struct
import zipfile
import zlib

PNG_1PX = (b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x02"
           b"\x00\x00\x00\x90wS\xde\x00\x00\x00\x0cIDATx\x9cc\xf8\xcf\xc0\x00\x00\x03\x01"
           b"\x01\x00\xc9\xfe\x92\xef\x00\x00\x00\x00IEND\xaeB`\x82")


def make_xlsx(path):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ventas"
    ws.append(["Producto", "Unidades", "Precio", "Total"])
    ws.append(["Café", 3, 2.5, "=B2*C2"])
    ws.append(["Té <b>", 2, 1.25, "=B3*C3"])
    ws["A1"].font = Font(bold=True, color="FF0000")
    ws["A1"].fill = PatternFill("solid", fgColor="FFFF00")
    ws.merge_cells("A5:B5")
    ws["A5"] = "Combinada"
    hidden = wb.create_sheet("Oculta")
    hidden.sheet_state = "hidden"
    wb.create_sheet("Resumen")["A1"] = 42
    wb.save(path)
    # Valores en caché de las fórmulas (como si lo hubiera guardado Excel)
    _set_cached_values(path, {"D2": "7.5", "D3": "2.5"})


def _set_cached_values(path, values):
    with zipfile.ZipFile(path) as zf:
        items = {n: zf.read(n) for n in zf.namelist()}
    xml = items["xl/worksheets/sheet1.xml"].decode()
    for ref, val in values.items():
        start = xml.index(f'<c r="{ref}"')
        end = xml.index("</c>", start)
        cell = xml[start:end]
        xml = xml[:start] + cell.replace("<v></v>", "").replace("</f>", f"</f><v>{val}</v>") + xml[end:]
    items["xl/worksheets/sheet1.xml"] = xml.encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        for n, d in items.items():
            zf.writestr(n, d)


def make_docx(path):
    import docx
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    d = docx.Document()
    d.add_heading("Informe trimestral", level=1)
    p = d.add_paragraph("Texto normal y ")
    p.add_run("negrita").bold = True
    p.add_run(" <script>alert(1)</script>")
    d.add_paragraph("Primer punto", style="List Bullet")
    d.add_paragraph("Segundo punto", style="List Bullet")
    d.add_paragraph("Paso uno", style="List Number")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "A"
    t.cell(0, 1).text = "B"
    t.cell(1, 0).text = "1"
    t.cell(1, 1).text = "2"
    # Enlace seguro y enlace javascript:
    for url, text in (("https://katphi.com", "web"), ("javascript:alert(1)", "malo")):
        para = d.add_paragraph()
        rid = para.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
        h = OxmlElement("w:hyperlink")
        h.set(qn("r:id"), rid)
        r = OxmlElement("w:r")
        tt = OxmlElement("w:t")
        tt.text = text
        r.append(tt)
        h.append(r)
        para._p.append(h)
    d.add_picture(io.BytesIO(PNG_1PX))
    d.save(path)


def make_pptx(path):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Katphi Office"
    s1.placeholders[1].text = "Subtítulo"
    s1.notes_slide.notes_text_frame.text = "Notas del orador"
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s2.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = "Caja <i>de texto</i>"
    run.font.size = Pt(28)
    run.font.bold = True
    s2.shapes.add_picture(io.BytesIO(PNG_1PX), Inches(5), Inches(2), Inches(1), Inches(1))
    tbl = s2.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
    tbl.cell(0, 0).text = "x"
    tbl.cell(1, 1).text = "y"
    prs.save(path)


def _odf(path, mime, content_xml):
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr(zipfile.ZipInfo("mimetype"), mime)
        zf.writestr("content.xml", content_xml)
        zf.writestr("META-INF/manifest.xml", "<manifest/>")


_ODF_HEAD = ('<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
             'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
             'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
             'xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0" '
             'xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0" '
             'xmlns:xlink="http://www.w3.org/1999/xlink">')


def make_odt(path):
    _odf(path, "application/vnd.oasis.opendocument.text", _ODF_HEAD +
         '<office:automatic-styles><style:style style:name="T1"><style:text-properties '
         'fo:font-weight="bold"/></style:style></office:automatic-styles>'
         '<office:body><office:text><text:h text:outline-level="2">Título ODT</text:h>'
         '<text:p>Hola <text:span text:style-name="T1">mundo</text:span> &lt;b&gt;</text:p>'
         '<text:list><text:list-item><text:p>item</text:p></text:list-item></text:list>'
         '</office:text></office:body></office:document-content>')


def make_ods(path):
    _odf(path, "application/vnd.oasis.opendocument.spreadsheet", _ODF_HEAD +
         '<office:body><office:spreadsheet><table:table table:name="Datos">'
         '<table:table-row><table:table-cell office:value-type="string"><text:p>a</text:p></table:table-cell>'
         '<table:table-cell office:value-type="float" office:value="5"><text:p>5</text:p></table:table-cell>'
         '<table:table-cell table:number-columns-repeated="16000"/></table:table-row>'
         '<table:table-row table:number-rows-repeated="1048000"><table:table-cell table:number-columns-repeated="1024"/></table:table-row>'
         '</table:table></office:spreadsheet></office:body></office:document-content>')


def make_xxe_odt(path):
    _odf(path, "application/vnd.oasis.opendocument.text",
         '<?xml version="1.0"?><!DOCTYPE x [<!ENTITY e SYSTEM "file:///etc/passwd">]>' +
         _ODF_HEAD + '<office:body><office:text><text:p>&e;</text:p></office:text></office:body>'
         '</office:document-content>')


def make_pdf(path, pages=2, text="Hola Katphi"):
    """PDF mínimo con texto (Helvetica) en cada página."""
    objs = []
    kids = " ".join(f"{3 + i * 2} 0 R" for i in range(pages))
    objs.append("<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {pages} >>")
    font_id = 3 + pages * 2
    for i in range(pages):
        content = f"BT /F1 24 Tf 72 720 Td ({text} {i + 1}) Tj ET".encode()
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                    f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {4 + i * 2} 0 R >>")
        objs.append(("stream", content))
    objs.append("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for n, obj in enumerate(objs, 1):
        offsets.append(out.tell())
        if isinstance(obj, tuple):
            data = obj[1]
            out.write(f"{n} 0 obj\n<< /Length {len(data)} >>\nstream\n".encode() + data + b"\nendstream\nendobj\n")
        else:
            out.write(f"{n} 0 obj\n{obj}\nendobj\n".encode())
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    with open(path, "wb") as f:
        f.write(out.getvalue())


def make_zip_bomb(path):
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("word/document.xml", b"\0" * (60 * 1024 * 1024))


def make_traversal_zip(path):
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr("word/../../evil.txt", "x")
