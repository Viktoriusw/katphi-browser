#!/usr/bin/env python3
"""
Herramientas de plugins premium para el asistente del chat.

El LLM usa los plugins Scraping, SEO Analyzer y Katphi Office escribiendo
etiquetas en su respuesta, igual que <search> o <nav>:

  <scrape>patterns</scrape>        detecta listas/tarjetas repetidas en la página
  <scrape>pattern 2</scrape>       selecciona los elementos del patrón 2
  <scrape>select .product</scrape> selecciona los elementos de un selector CSS
  <scrape>show</scrape> / <scrape>clear</scrape>
  <scrape>export csv|excel|json</scrape>
  <seo>analyze</seo> / <seo>keywords</seo>
  <seo>report pdf|docx|html|xlsx|md|json|csv</seo>  informe guardado en Descargas y abierto
  <office>sheet scraping|seo</office>, <office>sheet Título + tabla markdown</office>
  <office>read</office>             lee el documento abierto en la pestaña actual

Solo se ofrecen las herramientas de los plugins cargados; el gestor de plugins
únicamente carga un premium con licencia válida, así que el asistente no puede
saltarse la licencia. Los plugins no se modifican: se usan sus paneles y su
integración tal como los crea main_window.

Cada acción termina llamando a done(ok, aviso, resultado):
- aviso: frase corta para el chat (y la voz).
- resultado: texto para devolver al LLM y que continúe la respuesta, o "" si
  la acción ya terminó el turno (exportaciones, hojas creadas).
"""
from __future__ import annotations

import json
import logging
import os
import re
from datetime import datetime
from typing import Callable, List, Optional

from PySide6.QtCore import QObject, QThread, QTimer, Signal

logger = logging.getLogger(__name__)

Done = Callable[[bool, str, str], None]

TAGS = ("scrape", "seo", "office")
PLUGIN_IDS = {"scrape": "scraping", "seo": "seo_analyzer", "office": "office_suite"}
PLUGIN_NAMES = {"scrape": "Scraping", "seo": "SEO Analyzer", "office": "Katphi Office"}
# Órdenes que puede escribir el LLM (cada una es un método _<etiqueta>_<orden>)
COMMANDS = {
    "scrape": ("patterns", "pattern", "select", "show", "clear", "export"),
    "seo": ("analyze", "keywords", "report", "export"),
    "office": ("sheet", "read"),
}

HTML_TIMEOUT_MS = 8000
# Formatos del informe SEO (plugin 1.1.0) y alias que puede escribir el LLM
SEO_REPORT_FORMATS = ("pdf", "docx", "html", "xlsx", "md", "json", "csv")
_FORMAT_ALIASES = {"word": "docx", "doc": "docx", "excel": "xlsx", "xls": "xlsx", "markdown": "md",
                   "htm": "html", "web": "html"}
MAX_ROWS_FOR_LLM = 15
MAX_DOC_CHARS = 8000

PROMPT_HEADER = """

Premium plugins you can use (the user has them installed). Write the tag and one short sentence, then wait: the browser runs the plugin and sends you the result. Use them only when the user asks for something these plugins do, one tag per reply."""

PROMPT_SECTIONS = {
    "scrape": """
Scraping plugin (extract data from the current page):
- <scrape>patterns</scrape> detects repeated lists, cards and tables (products, results, articles...) and returns them numbered with their CSS selector.
- <scrape>pattern N</scrape> selects every element of pattern N; <scrape>select CSS</scrape> selects the elements of a CSS selector. Both return the selected rows.
- <scrape>show</scrape> shows the current selection; <scrape>clear</scrape> empties it.
- <scrape>export csv</scrape> (or excel, json) saves the selection to a file; the user chooses where.
- To extract data, first use <scrape>patterns</scrape>, choose the pattern that matches what the user wants, then <scrape>pattern N</scrape>. Never invent data: report only what the plugin returns.""",
    "seo": """
SEO Analyzer plugin (audits the current page):
- <seo>analyze</seo> analyzes the current page and returns the score, problems, recommendations and main keywords.
- <seo>keywords</seo> returns the page's keywords with their counts and density.
- <seo>report pdf</seo> creates the full SEO report as a file and opens it for the user. Formats: pdf, docx (Word), html, xlsx (Excel), md (Markdown), json, csv. It analyzes the page first if needed, so when the user asks for a report you can use it directly.
- After an analysis, explain the most important problems first and how to fix them.""",
    "office": """
Katphi Office plugin (spreadsheets and documents):
- <office>sheet scraping</office> creates a spreadsheet with the scraped selection and opens it in Katphi Office.
- <office>sheet seo</office> creates a spreadsheet with the latest SEO analysis (scores, problems, keywords).
- To create a spreadsheet with your own data (a comparison, a list, a plan...), put a title and a markdown table inside the tag:
<office>sheet Comparativa de portátiles
| Modelo | Precio | Peso |
|---|---|---|
| A | 999 | 1,2 |
</office>
- <office>read</office> reads the spreadsheet, document or PDF open in the current tab so you can summarize or analyze it.""",
}


class _CallThread(QThread):
    """Ejecuta una función lenta (análisis SEO) fuera del hilo de la interfaz."""
    result = Signal(object, str)   # valor, error

    def __init__(self, fn, parent=None):
        super().__init__(parent)
        self._fn = fn

    def run(self):
        try:
            self.result.emit(self._fn(), "")
        except Exception as e:  # el error se muestra en el chat
            logger.exception("Error en la herramienta de plugin")
            self.result.emit(None, str(e))


def _short(text, limit: int = 90) -> str:
    text = re.sub(r"\s+", " ", str(text or "")).strip()
    return text if len(text) <= limit else text[:limit - 1] + "…"


_NUMBER_RE = re.compile(r"^-?(?:\d{1,3}(?:\.\d{3})+|\d+)(?:,\d+)?$|^-?\d+(?:\.\d+)?$")


def cell_value(text: str):
    """Convierte '1.234,5' o '12.5' en número; el resto se queda como texto."""
    text = (text or "").strip()
    if not text or not _NUMBER_RE.match(text):
        return text
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    elif re.match(r"^-?\d{1,3}(?:\.\d{3})+$", text):
        text = text.replace(".", "")   # 1.234 = mil doscientos treinta y cuatro
    try:
        number = float(text)
    except ValueError:
        return text
    return int(number) if number.is_integer() and "." not in text else number


def parse_markdown_table(text: str):
    """Título (primera línea sin '|') y filas de una tabla markdown."""
    title, rows = "", []
    for line in (text or "").splitlines():
        line = line.strip()
        if not line:
            continue
        if not line.startswith("|"):
            if not rows and not title:
                title = line
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if all(re.match(r"^:?-{2,}:?$", c) for c in cells if c) and any(cells):
            continue   # separador |---|---|
        rows.append(cells)
    if rows:
        width = max(len(r) for r in rows)
        rows = [r + [""] * (width - len(r)) for r in rows]
    return title, rows


def effective_pattern_selector(pattern: dict) -> str:
    """Mismo criterio que el botón «Extraer patrón» del panel de scraping."""
    generic = {"div", "section", "main", "article", "ul", "ol", "table", "tbody", "tr"}
    sel = (pattern or {}).get("selectors") or {}
    rec = sel.get("recommended") or ""
    ancestor = sel.get("common_ancestor") or ""
    idx = sel.get("index_selectors") or []
    if rec and rec in generic and ancestor:
        return f"{ancestor} > {rec}"
    return rec or (idx[0] if idx else "")


def _pattern_size(p: dict) -> int:
    return p.get("elements_count") or p.get("items_count") or p.get("rows_count") or p.get("count") or 0


def _pattern_samples(p: dict, soup=None) -> List[str]:
    """Dos ejemplos de texto del patrón. Con la página (soup) se leen separando los
    elementos ("Modelo 1 · 19,95 €"); los del detector vienen pegados ("Modelo 119,95 €")."""
    if soup is not None:
        try:
            found = soup.select(effective_pattern_selector(p), limit=2)
            texts = [el.get_text(" · ", strip=True) for el in found]
            if any(texts):
                return texts
        except Exception:
            pass
    if p.get("sample_rows"):
        return [" | ".join(r.get("cell_contents", [])[:4]) for r in p["sample_rows"][:2]]
    samples = p.get("sample_items") or p.get("sample_elements") or []
    return [s.get("text", "") for s in samples[:2]]


class PluginTools(QObject):
    """Ejecuta las etiquetas <scrape>, <seo> y <office> del asistente."""

    def __init__(self, get_main_window, parent=None):
        super().__init__(parent)
        self._get_main_window = get_main_window
        self._patterns: list = []
        self._threads: list = []

    # ── Disponibilidad ─────────────────────────────────────────────────────

    @property
    def _main(self):
        return self._get_main_window()

    def _plugin_loaded(self, plugin_id: str) -> bool:
        pm = getattr(self._main, "plugin_manager", None)
        return bool(pm is not None and plugin_id in getattr(pm, "plugins", {}))

    def available(self) -> List[str]:
        """Etiquetas cuyos plugins están cargados (y por tanto con licencia)."""
        return [tag for tag in TAGS if self._plugin_loaded(PLUGIN_IDS[tag])]

    def prompt_section(self) -> str:
        tags = self.available()
        if not tags:
            return ""
        return PROMPT_HEADER + "".join(PROMPT_SECTIONS[t] for t in tags)

    # ── Ejecución ──────────────────────────────────────────────────────────

    def run(self, tag: str, arg: str, done: Done) -> None:
        tag = (tag or "").lower()
        if tag not in TAGS:
            done(False, f"Herramienta desconocida: {tag}", "")
            return
        if not self._plugin_loaded(PLUGIN_IDS[tag]):
            name = PLUGIN_NAMES[tag]
            done(False, f"El plugin {name} no está instalado o no tiene licencia activa.",
                 f"[Navegador] El plugin {name} no está disponible (no instalado o sin licencia). "
                 "Díselo al usuario; puede conseguirlo en la Plugin Store.")
            return
        raw = (arg or "").strip()
        command, _, rest = raw.partition("\n") if tag == "office" else (raw, "", "")
        verb, _, value = command.strip().partition(" ")
        verb = verb.lower()
        handler = getattr(self, f"_{tag}_{verb}", None) if verb in COMMANDS[tag] else None
        if handler is None:
            done(False, f"Orden no válida para {PLUGIN_NAMES[tag]}: {_short(command, 40)}",
                 f"[Navegador] La orden <{tag}>{_short(command, 40)}</{tag}> no existe. "
                 "Usa solo las órdenes indicadas.")
            return
        try:
            if tag == "office" and verb == "sheet":
                handler(value.strip(), rest, done)
            else:
                handler(value.strip(), done)
        except Exception as e:  # nunca romper el chat por un plugin
            logger.exception("Error ejecutando <%s>%s", tag, raw[:80])
            done(False, f"Error en {PLUGIN_NAMES[tag]}: {e}",
                 f"[Navegador] El plugin falló: {e}. Explícaselo al usuario.")

    # ── Utilidades ─────────────────────────────────────────────────────────

    def _current_web_tab(self):
        tm = getattr(self._main, "tab_manager", None)
        tabs = getattr(tm, "tabs", None) if tm else None
        browser = tabs.currentWidget() if tabs is not None else None
        if browser is None or not hasattr(browser, "page"):
            return None
        if browser.url().scheme() in ("katphi-app", "about", "") or browser.url().isEmpty():
            return None
        return browser

    def _with_page_html(self, done: Done, then: Callable[[object, str, str], None]) -> None:
        """Lee el HTML actual de la pestaña (DOM vivo) y llama a then(browser, html, url)."""
        browser = self._current_web_tab()
        if browser is None:
            done(False, "Abre una página web en la pestaña actual para usar el plugin.",
                 "[Navegador] La pestaña actual no es una página web. Pide al usuario que abra "
                 "la página que quiere analizar.")
            return
        url = browser.url().toString()
        state = {"called": False}

        def on_html(html):
            if state["called"]:
                return
            state["called"] = True
            if not html:
                done(False, "No he podido leer esta página.",
                     "[Navegador] No se pudo leer el HTML de la página.")
                return
            try:
                then(browser, html, url)
            except Exception as e:
                logger.exception("Error procesando el HTML para un plugin")
                done(False, f"Error en el plugin: {e}", f"[Navegador] El plugin falló: {e}.")

        browser.page().toHtml(on_html)
        QTimer.singleShot(HTML_TIMEOUT_MS, lambda: on_html(None))

    def _show_panel(self, attr: str, dock_attr: str) -> None:
        """Muestra el panel del plugin si vive en su dock. Si está en la pila lateral
        (donde también está el chat) no se cambia: taparía la conversación."""
        main = self._main
        panel = getattr(main, attr, None)
        dock = getattr(main, dock_attr, None)
        try:
            if panel is not None and dock is not None and panel.parentWidget() is dock \
                    and not dock.isVisible():
                dock.show()
                dock.raise_()
        except Exception as e:
            logger.debug("No se pudo mostrar el panel %s: %s", attr, e)

    def _run_in_thread(self, fn, on_result) -> None:
        thread = _CallThread(fn, self)
        self._threads.append(thread)
        thread.result.connect(on_result)
        thread.finished.connect(lambda t=thread: self._threads.remove(t) if t in self._threads else None)
        thread.start()

    def shutdown(self, timeout_ms: int = 3000) -> None:
        for thread in list(self._threads):
            thread.wait(timeout_ms)

    # ── Scraping ───────────────────────────────────────────────────────────

    def _scraping(self):
        main = self._main
        integration = getattr(main, "scraping_integration", None)
        panel = getattr(main, "scraping_panel", None)
        if integration is None or panel is None:
            raise RuntimeError("el panel de Scraping no se ha cargado")
        return integration, panel

    def _scraping_load(self, browser, html: str, url: str):
        integration, panel = self._scraping()
        integration.update_content(html, url)
        integration.browser_widget = browser
        panel.browser_tab = browser
        return integration, panel

    def _selection_feedback(self, integration, header: str) -> str:
        rows = integration._build_export_rows() if integration.get_selected_elements() else []
        if not rows:
            return header + "\nLa selección está vacía."
        columns = [c for c in ("titulo", "precio", "texto", "enlace") if c in rows[0]]
        lines = [f"{header}\nSelección actual: {len(rows)} elementos. "
                 f"Primeros {min(len(rows), MAX_ROWS_FOR_LLM)}:"]
        for r in rows[:MAX_ROWS_FOR_LLM]:
            lines.append(f"{r['n']}. " + " | ".join(
                f"{c}: {_short(r[c], 80 if c != 'enlace' else 120)}" for c in columns if r.get(c)))
        return "\n".join(lines)

    def _scrape_patterns(self, _value: str, done: Done) -> None:
        def then(browser, html, url):
            integration, panel = self._scraping_load(browser, html, url)
            if not hasattr(integration, "detect_dom_patterns"):
                done(False, "El detector de patrones del plugin no está disponible.",
                     "[Navegador] El detector de patrones no está disponible.")
                return
            result = integration.detect_dom_patterns()
            if "error" in result:
                done(False, f"Scraping: {result['error']}", f"[Navegador] Error: {result['error']}")
                return
            self._patterns = [p for p in result.get("patterns", []) if effective_pattern_selector(p)]
            try:
                panel.detected_patterns = self._patterns
                panel.display_patterns(self._patterns)
            except Exception:
                pass
            self._show_panel("scraping_panel", "scraping_dock")
            if not self._patterns:
                done(True, "No he encontrado listas ni tarjetas repetidas en esta página.",
                     "[Plugin Scraping] No hay patrones repetidos en la página. Propón al usuario "
                     "un selector CSS concreto o decirle qué datos quiere.")
                return
            from bs4 import BeautifulSoup
            soup = BeautifulSoup(html, "html.parser")
            lines = [f"[Plugin Scraping] Patrones detectados en {url}:"]
            for i, p in enumerate(self._patterns[:10], 1):
                samples = "; ".join(_short(s, 80) for s in _pattern_samples(p, soup) if s)
                lines.append(f"{i}. {p.get('type', '?')} — {_pattern_size(p)} elementos — "
                             f"selector {effective_pattern_selector(p)}"
                             + (f" — ejemplos: {samples}" if samples else ""))
            lines.append("Elige el patrón que corresponde a lo que pidió el usuario y usa "
                         "<scrape>pattern N</scrape>, o pregúntale si no está claro.")
            done(True, f"Scraping: {len(self._patterns)} patrones detectados.", "\n".join(lines))
        self._with_page_html(done, then)

    def _scrape_pattern(self, value: str, done: Done) -> None:
        if not value.isdigit() or not 1 <= int(value) <= len(self._patterns):
            done(False, "Ese patrón no existe; primero hay que detectar los patrones.",
                 f"[Navegador] No hay patrón {value}. Usa antes <scrape>patterns</scrape>.")
            return
        self._scrape_select(effective_pattern_selector(self._patterns[int(value) - 1]), done)

    def _scrape_select(self, selector: str, done: Done) -> None:
        selector = selector.strip().strip("`'\"")
        if not selector:
            done(False, "Falta el selector CSS.", "[Navegador] Falta el selector CSS.")
            return

        def then(browser, html, url):
            integration, panel = self._scraping_load(browser, html, url)
            result = integration.add_element_by_selector(selector)
            if "error" in result:
                done(False, f"Scraping: ningún elemento coincide con {selector}",
                     f"[Plugin Scraping] {result['error']}. Prueba otro patrón o selector.")
                return
            try:
                panel.update_selected_list()
            except Exception:
                pass
            self._show_panel("scraping_panel", "scraping_dock")
            added = result.get("elements_added", 0)
            done(True, f"Scraping: {added} elementos añadidos a la selección.",
                 self._selection_feedback(integration, f"[Plugin Scraping] Añadidos {added} "
                                                       f"elementos con el selector {selector}."))
        self._with_page_html(done, then)

    def _scrape_show(self, _value: str, done: Done) -> None:
        integration, _panel = self._scraping()
        count = len(integration.get_selected_elements())
        done(True, f"Scraping: {count} elementos seleccionados.",
             self._selection_feedback(integration, "[Plugin Scraping]"))

    def _scrape_clear(self, _value: str, done: Done) -> None:
        integration, panel = self._scraping()
        integration.clear_selected_elements()
        try:
            panel.update_selected_list()
        except Exception:
            pass
        done(True, "Scraping: selección vaciada.", "")

    def _scrape_export(self, value: str, done: Done) -> None:
        integration, panel = self._scraping()
        fmt = {"csv": "CSV", "excel": "Excel", "xlsx": "Excel", "json": "JSON",
               "yaml": "YAML"}.get(value.lower() or "csv")
        if fmt is None:
            done(False, f"Formato no válido: {value}", "")
            return
        if not integration.get_selected_elements():
            done(False, "No hay nada seleccionado para exportar.",
                 "[Navegador] La selección está vacía; primero hay que seleccionar datos.")
            return
        self._show_panel("scraping_panel", "scraping_dock")
        panel.export_format.setCurrentText(fmt)
        panel.run_export()   # el usuario elige dónde guardar
        done(True, f"Scraping: exportación a {fmt} lanzada.", "")

    # ── SEO ────────────────────────────────────────────────────────────────

    def _seo_panel(self):
        panel = getattr(self._main, "seo_panel", None)
        if panel is None or not hasattr(panel, "coordinator"):
            raise RuntimeError("el panel de SEO Analyzer no se ha cargado")
        return panel

    def _seo_analyze(self, _value: str, done: Done, keywords_only: bool = False,
                     after: Optional[Callable[[], None]] = None) -> None:
        """Analiza la página actual. Con `after`, al terminar se llama a after() en vez de done."""
        panel = self._seo_panel()

        def then(browser, html, url):
            coordinator = panel.coordinator

            def on_result(results, error):
                if error or not results:
                    done(False, f"SEO: error en el análisis: {error}",
                         f"[Plugin SEO] El análisis falló: {error}.")
                    return
                panel.current_url = url
                panel.current_html = html
                try:
                    panel.url_input.setText(url)
                except Exception:
                    pass
                panel.on_analysis_finished(results)   # pinta el panel y calcula las keywords
                self._show_panel("seo_panel", "seo_dock")
                if after is not None:
                    after()
                    return
                score = results.get("overall", {}).get("score", 0)
                feedback = (self._seo_keywords_text(panel, url, 40) if keywords_only
                            else self._seo_summary_text(panel, results, url))
                done(True, f"SEO: análisis completado, puntuación {score}/100.", feedback)

            self._run_in_thread(lambda: coordinator.analyze(html, url), on_result)
        self._with_page_html(done, then)

    def _seo_keywords(self, value: str, done: Done) -> None:
        panel = self._seo_panel()
        browser = self._current_web_tab()
        url = browser.url().toString() if browser is not None else ""
        if panel.current_keywords and url and panel.current_url == url:
            done(True, "SEO: palabras clave de la página.", self._seo_keywords_text(panel, url, 40))
            return
        self._seo_analyze(value, done, keywords_only=True)

    @staticmethod
    def _seo_keywords_text(panel, url: str, limit: int) -> str:
        keywords = list(getattr(panel, "current_keywords", []) or [])
        if not keywords:
            return f"[Plugin SEO] No se han encontrado palabras clave repetidas en {url}."
        lines = [f"[Plugin SEO] Palabras clave de {url} (palabra/frase — apariciones — "
                 "densidad — en título/H1/description):"]
        for k in keywords[:limit]:
            lines.append(f"- {k['keyword']} — {k['count']} — {k['density']} % — "
                         f"{'sí' if k.get('in_key_tags') else 'no'}")
        return "\n".join(lines)

    def _seo_summary_text(self, panel, results: dict, url: str) -> str:
        coordinator = panel.coordinator
        overall = results.get("overall", {})
        lines = [f"[Plugin SEO] Análisis de {url}: puntuación global {overall.get('score', 0)}/100."]
        try:
            summary = coordinator.get_summary()
            lines.append(f"Problemas críticos: {summary.get('critical_issues', 0)}, avisos: "
                         f"{summary.get('warnings', 0)}, info: {summary.get('info', 0)}.")
        except Exception:
            pass
        for name, res in results.items():
            if name == "overall" or not isinstance(res, dict):
                continue
            try:
                label = coordinator.get_analyzer_info(name).get("name", name)
            except Exception:
                label = name
            lines.append(f"\n{label}: {res.get('score', 0)}/100")
            for issue in (res.get("issues") or [])[:5]:
                lines.append(f"  - [{issue.get('type', '')}] {_short(issue.get('message', ''), 140)}")
        try:
            recs = coordinator.get_recommendations()[:8]
        except Exception:
            recs = []
        if recs:
            lines.append("\nRecomendaciones:")
            lines += [f"- {_short(r, 160)}" for r in recs]
        top = [k["keyword"] for k in (panel.current_keywords or [])[:12]]
        if top:
            lines.append("\nPalabras clave principales: " + ", ".join(top))
        lines.append("\nEl análisis completo está en el panel SEO. Si el usuario pidió un informe o un archivo "
                     "(PDF, Word, Excel...), escribe ahora <seo>report FORMATO</seo>; si no, resume lo más "
                     "importante para el usuario.")
        return "\n".join(lines)

    def _seo_report(self, value: str, done: Done) -> None:
        panel = self._seo_panel()
        fmt = (value or "pdf").lower().strip(". ")
        fmt = _FORMAT_ALIASES.get(fmt, fmt)
        if fmt not in SEO_REPORT_FORMATS:
            done(False, f"Formato de informe no válido: {value}",
                 f"[Navegador] Formato {value} no soportado. Usa: {', '.join(SEO_REPORT_FORMATS)}.")
            return
        if not hasattr(panel, "save_report"):
            self._seo_report_legacy(panel, fmt, done)
            return
        browser = self._current_web_tab()
        url = browser.url().toString() if browser is not None else ""
        if panel.current_results and (not url or panel.current_url == url):
            self._write_seo_report(panel, fmt, done)
        else:
            # Sin análisis de esta página: se analiza y después se crea el informe
            self._seo_analyze("", done, after=lambda: self._write_seo_report(panel, fmt, done))

    _seo_export = _seo_report

    def _write_seo_report(self, panel, fmt: str, done: Done) -> None:
        from urllib.parse import urlparse
        host = re.sub(r"[^\w.-]+", "_", urlparse(panel.current_url).netloc or "pagina")[:50]
        path = self._unique_path(f"informe_seo_{host}", fmt)
        try:
            path = panel.save_report(fmt, path)
        except Exception as e:
            logger.exception("No se pudo crear el informe SEO")
            done(False, f"No se pudo crear el informe: {e}",
                 f"[Navegador] El informe SEO falló: {e}. Explícaselo al usuario.")
            return
        logger.info("Informe SEO creado por el asistente: %s", path)
        opened = self._open_file(path, fmt)
        name = os.path.basename(path)
        where = "y abierto" if opened else f"en {os.path.dirname(path)}"
        done(True, f"Informe SEO «{name}» guardado {where}.", "")

    def _seo_report_legacy(self, panel, fmt: str, done: Done) -> None:
        """SEO Analyzer 1.0.x: solo exporta con su diálogo (json, csv, pdf con ReportLab)."""
        if fmt not in ("pdf", "json", "csv"):
            done(False, "Tu versión del plugin SEO solo exporta PDF, JSON y CSV; actualízalo desde la Plugin Store.", "")
            return
        if not panel.current_results:
            done(False, "Primero hay que analizar la página.",
                 "[Navegador] No hay análisis SEO todavía; usa antes <seo>analyze</seo>.")
            return
        self._show_panel("seo_panel", "seo_dock")
        panel.export_results(fmt)   # el usuario elige dónde guardar
        done(True, f"SEO: exportación a {fmt.upper()} lanzada.", "")

    def _open_file(self, path: str, fmt: str) -> bool:
        """Entrega el archivo al usuario: HTML en una pestaña, documentos en Katphi Office."""
        try:
            if fmt == "html":
                from PySide6.QtCore import QUrl
                tm = getattr(self._main, "tab_manager", None)
                if tm is None:
                    return False
                tm.add_new_tab(QUrl.fromLocalFile(path).toString(), focus=True)
                return True
            if fmt in ("pdf", "docx", "xlsx", "md", "csv"):
                return self._open_in_office(path)
        except Exception as e:
            logger.warning("No se pudo abrir %s: %s", path, e)
        return False

    # ── Katphi Office ──────────────────────────────────────────────────────

    def _office_sheet(self, source: str, table: str, done: Done) -> None:
        source_key = source.lower()
        if source_key == "scraping":
            if not self._plugin_loaded("scraping"):
                done(False, "Para esto hace falta el plugin Scraping.", "")
                return
            integration, _panel = self._scraping()
            if not integration.get_selected_elements():
                done(False, "No hay datos de scraping seleccionados.",
                     "[Navegador] La selección de scraping está vacía; primero selecciona datos.")
                return
            rows = integration._build_export_rows()
            sheets = [("Datos", list(rows[0].keys()), [list(r.values()) for r in rows])]
            title = "scraping"
        elif source_key == "seo":
            if not self._plugin_loaded("seo_analyzer"):
                done(False, "Para esto hace falta el plugin SEO Analyzer.", "")
                return
            panel = self._seo_panel()
            if not panel.current_results:
                done(False, "Primero hay que analizar la página con SEO Analyzer.",
                     "[Navegador] No hay análisis SEO todavía; usa antes <seo>analyze</seo>.")
                return
            sheets = self._seo_sheets(panel)
            title = "seo"
        else:
            title, rows = parse_markdown_table(f"{source}\n{table}")
            if len(rows) < 2:
                done(False, "No he recibido ninguna tabla para la hoja.",
                     "[Navegador] Para crear una hoja con tus datos escribe una tabla markdown "
                     "dentro de <office>sheet Título ...</office>.")
                return
            sheets = [("Hoja1", rows[0], [[cell_value(c) for c in r] for r in rows[1:]])]
            title = title or "tabla"

        path = self._write_workbook(title, sheets)
        opened = self._open_in_office(path)
        total = sum(len(s[2]) for s in sheets)
        name = os.path.basename(path)
        if opened:
            done(True, f"Katphi Office: hoja «{name}» creada con {total} filas y abierta.", "")
        else:
            done(True, f"Hoja guardada en {path} ({total} filas); no se pudo abrir en Katphi Office.", "")

    @staticmethod
    def _seo_sheets(panel) -> list:
        results = panel.current_results
        coordinator = panel.coordinator

        def label(name):
            try:
                return coordinator.get_analyzer_info(name).get("name", name)
            except Exception:
                return name

        summary, issues = [], []
        for name, res in results.items():
            if name == "overall" or not isinstance(res, dict):
                continue
            summary.append([label(name), res.get("score", 0), res.get("severity", "")])
            for issue in res.get("issues") or []:
                issues.append([label(name), issue.get("type", ""), issue.get("message", ""),
                               _short(issue.get("element", ""), 300)])
        summary.append(["Global", results.get("overall", {}).get("score", 0), panel.current_url])
        keywords = [[k["keyword"], k["type"], k["count"], k["density"],
                     "sí" if k.get("in_key_tags") else "no"]
                    for k in (panel.current_keywords or [])]
        sheets = [("Resumen", ["Análisis", "Puntuación", "Estado"], summary),
                  ("Problemas", ["Análisis", "Tipo", "Mensaje", "Elemento"], issues)]
        if keywords:
            sheets.append(("Palabras clave", ["Palabra clave", "Tipo", "Apariciones",
                                              "Densidad %", "En título/H1/description"], keywords))
        return sheets

    @staticmethod
    def _output_dir() -> str:
        try:
            from constants import is_private_mode
            private = bool(is_private_mode())
        except Exception:
            private = False
        if private:
            import tempfile
            return tempfile.mkdtemp(prefix="katphi-office-")
        from PySide6.QtCore import QStandardPaths
        folder = QStandardPaths.writableLocation(QStandardPaths.DownloadLocation) \
            or os.path.expanduser("~")
        os.makedirs(folder, exist_ok=True)
        return folder

    def _unique_path(self, name: str, ext: str) -> str:
        """Ruta nueva en Descargas (nunca sobrescribe un archivo existente)."""
        base = os.path.join(self._output_dir(), f"{name}_{datetime.now():%Y%m%d_%H%M%S}")
        path, n = f"{base}.{ext}", 1
        while os.path.exists(path):
            n += 1
            path = f"{base}_{n}.{ext}"
        return path

    def _write_workbook(self, title: str, sheets: list) -> str:
        """Crea un .xlsx nuevo (nunca sobrescribe) con cabecera en negrita y filtro."""
        import openpyxl
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for sheet_name, header, rows in sheets:
            ws = wb.create_sheet(re.sub(r"[\[\]:*?/\\]", "", sheet_name)[:31] or "Hoja")
            ws.append([str(h) for h in header])
            for row in rows:
                ws.append([v if isinstance(v, (int, float)) else str(v or "") for v in row])
            for cell in ws[1]:
                cell.font = Font(bold=True)
            ws.freeze_panes = "A2"
            if rows:
                ws.auto_filter.ref = ws.dimensions
            for col in range(1, len(header) + 1):
                values = [str(ws.cell(row=r, column=col).value or "") for r in range(1, min(ws.max_row, 200) + 1)]
                ws.column_dimensions[get_column_letter(col)].width = min(max(map(len, values), default=8) + 2, 60)

        slug = re.sub(r"[^\w-]+", "_", title.lower(), flags=re.UNICODE).strip("_")[:40] or "tabla"
        path = self._unique_path(f"katphi_{slug}", "xlsx")
        wb.save(path)
        logger.info("Hoja creada por el asistente: %s", path)
        return path

    def _open_in_office(self, path: str) -> bool:
        try:
            from katphi_office import app
            return bool(app.open_document(self._main, path))
        except Exception as e:
            logger.warning("No se pudo abrir %s en Katphi Office: %s", path, e)
            return False

    def _office_read(self, _value: str, done: Done) -> None:
        try:
            from katphi_office.sessions import STORE
        except ImportError:
            done(False, "El visor de Katphi Office no está disponible.", "")
            return
        tm = getattr(self._main, "tab_manager", None)
        current = tm.tabs.currentWidget() if tm is not None else None
        session = next((s for s in STORE.sessions() if s.extra.get("browser") is current), None)
        if session is None:
            done(False, "La pestaña actual no es un documento de Katphi Office.",
                 "[Navegador] La pestaña actual no tiene un documento de Katphi Office abierto. "
                 "Pide al usuario que abra la hoja, el documento o el PDF en una pestaña.")
            return
        if session.extra.get("new") or not session.path:
            done(False, "Guarda el documento primero para que pueda leerlo.",
                 "[Navegador] El documento es nuevo y aún no se ha guardado; pide al usuario que lo guarde.")
            return
        if session.status != "ready":
            done(False, "El documento aún se está cargando.",
                 "[Navegador] El documento todavía se está cargando; pide al usuario que espere y repita.")
            return

        def read():
            return self._document_text(session)

        def on_result(text, error):
            if error or not text:
                done(False, f"No he podido leer «{session.name}».",
                     f"[Navegador] No se pudo leer el documento: {error or 'sin texto'}.")
                return
            note = ("" if not session.dirty else
                    "\n(Nota: el documento tiene cambios sin guardar que no se incluyen.)")
            done(True, f"Katphi Office: leyendo «{session.name}».",
                 f"[Plugin Katphi Office] Contenido de «{session.name}»:{note}\n{text}")

        self._run_in_thread(read, on_result)

    @staticmethod
    def _document_text(session) -> str:
        if session.family == "pdf":
            from pypdf import PdfReader
            reader = PdfReader(session.view_path)
            parts = []
            for i, page in enumerate(reader.pages[:30], 1):
                parts.append(f"--- Página {i} ---\n{page.extract_text() or ''}")
                if sum(map(len, parts)) > MAX_DOC_CHARS:
                    break
            return "\n".join(parts)[:MAX_DOC_CHARS]
        model = json.loads(session.model_json or b"{}")
        if session.family == "sheets":
            out = []
            for sheet in model.get("sheets", []):
                if sheet.get("hidden"):
                    continue
                out.append(f"--- Hoja «{sheet.get('name', '')}» ---")
                for row in sheet.get("rows", [])[:200]:
                    out.append("\t".join((c or {}).get("v", "") if isinstance(c, dict) else "" for c in row))
                if sum(map(len, out)) > MAX_DOC_CHARS:
                    break
            return "\n".join(out)[:MAX_DOC_CHARS]
        if session.family == "document":
            from bs4 import BeautifulSoup
            text = BeautifulSoup(model.get("html", ""), "html.parser").get_text("\n")
            return re.sub(r"\n{3,}", "\n\n", text).strip()[:MAX_DOC_CHARS]
        texts = []

        def walk(node):
            if isinstance(node, dict):
                for k, v in node.items():
                    if k == "text" and isinstance(v, str) and v.strip():
                        texts.append(v.strip())
                    else:
                        walk(v)
            elif isinstance(node, list):
                for v in node:
                    walk(v)
        walk(model)
        return "\n".join(texts)[:MAX_DOC_CHARS]
